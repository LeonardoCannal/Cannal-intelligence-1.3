from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for
from scraper import (
    buscar_medicos_doctoralia,
    buscar_medicos_sechat,
    buscar_medicos_fonte_generica,
    FONTES_GENERICAS,
)
from auth import (
    inicializar_auth,
    criar_usuario,
    buscar_usuario_por_email,
    buscar_usuario_por_cpf,
    buscar_usuario_por_id,
    cpf_valido,
    formatar_cpf,
    checar_senha,
    atualizar_ultimo_login,
    listar_usuarios_com_contagem_painel,
    usuario_logado,
    fazer_login,
    fazer_logout,
    login_required,
    admin_required,
    ErroIntegridade,
    computar_chave_medico,
    adicionar_ao_painel,
    mover_no_painel,
    listar_painel_usuario,
    chaves_no_painel_usuario,
)
import webbrowser
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
import sys
import os
import io
import secrets
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def resource_path(relative_path):
    """Acha o caminho certo dos arquivos SÓ DE LEITURA (templates, static),
    tanto rodando normal quanto dentro do .exe empacotado."""
    base_path = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base_path, relative_path)


def caminho_dados_persistentes(nome_arquivo):
    """
    Acha o caminho certo pra arquivos que precisam ser GRAVADOS e
    PERSISTIR entre execuções (como o banco de dados de usuários).
    Dentro do .exe, a pasta que o resource_path() usa é temporária e é
    apagada quando o programa fecha — então o banco de dados precisa
    ficar do lado do .exe, não dentro dela.
    """
    if getattr(sys, "frozen", False):
        base_path = os.path.dirname(sys.executable)
    else:
        base_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_path, nome_arquivo)


app = Flask(
    __name__,
    template_folder=resource_path("templates"),
    static_folder=resource_path("static"),
)

# A SECRET_KEY protege as sessões de login. Em produção (site no ar),
# defina a variável de ambiente SECRET_KEY com um valor fixo e secreto —
# senão, toda vez que o servidor reiniciar, todo mundo é deslogado.
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", secrets.token_hex(32))

inicializar_auth(app, caminho_dados_persistentes("cannal.db"))

MAX_ESPECIALIDADES_POR_BUSCA = 3

ESPECIALIDADES = {
    "neurologista": {"nome": "Neurologista", "doctoralia": "neurologista", "sechat": "n6k3wbvn8k5ajvtb75ikhaaw"},
    "psiquiatra": {"nome": "Psiquiatra", "doctoralia": "psiquiatra", "sechat": "jiagob2yvu3v78j041jb70oz"},
    "neurologista-pediatrico": {"nome": "Neuropediatra", "doctoralia": "neurologista-pediatrico", "sechat": "k9b3og0pgrbrizm138hkamr8"},
    "medico-clinico-geral": {"nome": "Clínico / Clínica Médica", "doctoralia": "medico-clinico-geral", "sechat": "eol7qqtdgtvf8buc2px9hbrq"},
    "pediatra": {"nome": "Pediatra", "doctoralia": "pediatra", "sechat": "dfwzf1l7z9efzucsufeqqeu9"},
    "medico-de-familia": {"nome": "Medicina de Família", "doctoralia": "medico-de-familia", "sechat": "yycwljzp344y327g6rkuyvp1"},
    "especialista-em-dor": {"nome": "Médico da Dor / Anestesiologia", "doctoralia": "especialista-em-dor", "sechat": "zyqcuzojsnjny7wl1lk1g7ah"},
    "ortopedista-traumatologista": {"nome": "Ortopedista", "doctoralia": "ortopedista-traumatologista", "sechat": "wc5z54hdmebdmv8n783876nk"},
    "oncologista": {"nome": "Oncologista", "doctoralia": "oncologista", "sechat": "r742v7goz2k4q6467wsk99di"},
    "geriatra": {"nome": "Geriatra", "doctoralia": "geriatra", "sechat": "ly7dotxntqq446jbqlg6b3l5"},
    "alergista": {"nome": "Alergologista", "doctoralia": "alergista", "sechat": "homqkrplnu5aqeu2uota9wu9"},
    "cardiologista": {"nome": "Cardiologista", "doctoralia": "cardiologista", "sechat": "x14ir5c9kvx9s11bwbvi1mys"},
    "dermatologista": {"nome": "Dermatologista", "doctoralia": "dermatologista", "sechat": "pylctsugyjl97vtfjjs8da1k"},
    "endocrinologista": {"nome": "Endocrinologista", "doctoralia": "endocrinologista", "sechat": "u28uuuaeimpxf3k25cnxru8o"},
    "endocrinologista-pediatrico": {"nome": "Endocrinologista Pediátrico", "doctoralia": "endocrinologista-pediatrico", "sechat": None},
    "especialista-em-medicina-fisica-e-reabilitacao": {"nome": "Medicina Física e Reabilitação", "doctoralia": "especialista-em-medicina-fisica-e-reabilitacao", "sechat": "alfpewwlicu8tux4xcta4yh5"},
    "especialista-em-medicina-preventiva": {"nome": "Medicina Preventiva", "doctoralia": "especialista-em-medicina-preventiva", "sechat": "wutli6r3320ylrt2uclsfp57"},
    "gastroenterologista": {"nome": "Gastroenterologista", "doctoralia": "gastroenterologista", "sechat": "f6w5gv0vbyw6tpambkg1mxag"},
    "generalista": {"nome": "Generalista", "doctoralia": "generalista", "sechat": "c6vsx97abcofv176xjliox2w"},
    "ginecologista": {"nome": "Ginecologista", "doctoralia": "ginecologista", "sechat": "mlhojg8vxi8iftl7jbgudfsw"},
    "medico-acupunturista": {"nome": "Médico Acupunturista", "doctoralia": "medico-acupunturista", "sechat": "djkdf9qhvdw4jzfbhgp4o1pk"},
    "medico-do-esporte": {"nome": "Médico do Esporte", "doctoralia": "medico-do-esporte", "sechat": "kqs3bhoscy5nnf3vdb3vxn1w"},
    "medico-do-sono": {"nome": "Médico do Sono", "doctoralia": "medico-do-sono", "sechat": None},
    "medico-do-trabalho": {"nome": "Médico do Trabalho", "doctoralia": "medico-do-trabalho", "sechat": "evc6v5o06an5xtefc9339ay9"},
    "traumatologista": {"nome": "Traumatologista", "doctoralia": "ortopedista-traumatologista", "sechat": "wc5z54hdmebdmv8n783876nk"},
    "reumatologista": {"nome": "Reumatologista", "doctoralia": "reumatologista", "sechat": "qeiljtcgscrqrwp7wy8jr3cp"},
}

NOMES_ESPECIALIDADES = {slug: info["nome"] for slug, info in ESPECIALIDADES.items()}


@app.route("/cadastro", methods=["GET", "POST"])
def cadastro():
    if usuario_logado():
        return redirect(url_for("index"))

    erro = None
    nome_form = ""
    email_form = ""
    cpf_form = ""

    if request.method == "POST":
        nome_form = request.form.get("nome", "").strip()
        email_form = request.form.get("email", "").strip().lower()
        cpf_form = request.form.get("cpf", "").strip()
        senha = request.form.get("senha", "")
        confirmar = request.form.get("confirmar", "")

        if not nome_form or not email_form or not cpf_form or not senha:
            erro = "Preencha todos os campos."
        elif not cpf_valido(cpf_form):
            erro = "CPF inválido. Confira os números digitados."
        elif senha != confirmar:
            erro = "As senhas não coincidem."
        elif len(senha) < 6:
            erro = "A senha precisa ter pelo menos 6 caracteres."
        elif buscar_usuario_por_email(email_form):
            erro = "Já existe uma conta cadastrada com esse e-mail."
        elif buscar_usuario_por_cpf(cpf_form):
            erro = "Já existe uma conta cadastrada com esse CPF."

        if not erro:
            try:
                user_id = criar_usuario(nome_form, email_form, cpf_form, senha)
            except ErroIntegridade:
                erro = "Já existe uma conta cadastrada com esse e-mail ou CPF."
            else:
                atualizar_ultimo_login(user_id)
                fazer_login(user_id)
                return redirect(url_for("index"))

    return render_template("cadastro.html", erro=erro, nome=nome_form, email=email_form, cpf=cpf_form)


@app.route("/login", methods=["GET", "POST"])
def login():
    if usuario_logado():
        return redirect(url_for("index"))

    erro = None
    if request.method == "POST":
        cpf = request.form.get("cpf", "").strip()
        senha = request.form.get("senha", "")
        usuario = buscar_usuario_por_cpf(cpf)

        if usuario and checar_senha(usuario, senha):
            atualizar_ultimo_login(usuario["id"])
            fazer_login(usuario["id"])
            return redirect(url_for("index"))

        erro = "CPF ou senha incorretos."

    return render_template("login.html", erro=erro)


@app.route("/logout")
@login_required
def logout():
    fazer_logout()
    return redirect(url_for("login"))


@app.route("/admin")
@admin_required
def admin():
    def formatar_data(valor_iso):
        if not valor_iso:
            return "Nunca"
        try:
            return datetime.fromisoformat(valor_iso).strftime("%d/%m/%Y %H:%M")
        except ValueError:
            return valor_iso

    usuarios = [
        {
            "id": u["id"],
            "nome": u["nome"],
            "email": u["email"],
            "cpf": formatar_cpf(u["cpf"]),
            "criado_em": formatar_data(u["criado_em"]),
            "ultimo_login": formatar_data(u["ultimo_login"]),
            "is_admin": bool(u["is_admin"]),
            "total_painel": u["total_painel"],
        }
        for u in listar_usuarios_com_contagem_painel()
    ]
    return render_template("admin.html", usuarios=usuarios)


@app.route("/admin/exportar/<int:user_id>")
@admin_required
def admin_exportar_painel(user_id):
    usuario = buscar_usuario_por_id(user_id)
    if not usuario:
        return "Usuário não encontrado.", 404

    painel = listar_painel_usuario(user_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Painel Médico"

    colunas = ["Nome", "CRM", "Especialidade", "Cidade", "UF", "Endereço", "Telefone", "Status", "Adicionado em"]
    ws.append(colunas)

    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill(start_color="343C4C", end_color="343C4C", fill_type="solid")
    for celula in ws[1]:
        celula.font = cabecalho_fonte
        celula.fill = cabecalho_fundo

    for m in painel:
        try:
            data_fmt = datetime.fromisoformat(m["adicionado_em"]).strftime("%d/%m/%Y %H:%M")
        except (ValueError, TypeError):
            data_fmt = m["adicionado_em"] or ""

        ws.append([
            m["nome"], m["crm"], m["especialidade"], m["cidade"], m["uf"],
            m["endereco"], m["telefone"],
            "Visitado" if m["status"] == "visitado" else "Prospectado",
            data_fmt,
        ])

    larguras = [28, 16, 24, 18, 6, 38, 18, 14, 18]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = largura

    arquivo = io.BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)

    nome_usuario_arquivo = usuario["nome"].strip().replace(" ", "_")
    nome_arquivo = f"painel_{nome_usuario_arquivo}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return send_file(
        arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=nome_arquivo,
    )


@app.route("/admin/painel/<int:user_id>")
@admin_required
def admin_painel_usuario(user_id):
    usuario = buscar_usuario_por_id(user_id)
    if not usuario:
        return jsonify({"erro": "Usuário não encontrado."}), 404

    painel = listar_painel_usuario(user_id)
    prospectados = [dict(m) for m in painel if m["status"] == "prospectado"]
    visitados = [dict(m) for m in painel if m["status"] == "visitado"]

    return jsonify({
        "usuario": {"nome": usuario["nome"], "email": usuario["email"]},
        "prospectados": prospectados,
        "visitados": visitados,
    })


@app.route("/")
@login_required
def index():
    return render_template(
        "index.html",
        especialidades=NOMES_ESPECIALIDADES,
        max_especialidades=MAX_ESPECIALIDADES_POR_BUSCA,
    )


@app.route("/api/buscar")
@login_required
def api_buscar():
    especialidades_selecionadas = request.args.getlist("especialidade")
    cidade = request.args.get("cidade", "")
    uf = request.args.get("uf", "")

    if not especialidades_selecionadas:
        return jsonify({"erro": "Selecione ao menos uma especialidade."}), 400
    if len(especialidades_selecionadas) > MAX_ESPECIALIDADES_POR_BUSCA:
        return jsonify({"erro": f"Selecione no máximo {MAX_ESPECIALIDADES_POR_BUSCA} especialidades por busca."}), 400
    for slug in especialidades_selecionadas:
        if slug not in ESPECIALIDADES:
            return jsonify({"erro": f"Especialidade inválida: {slug}"}), 400
    if not cidade.strip():
        return jsonify({"erro": "Informe a cidade."}), 400

    nomes_especialidades = [ESPECIALIDADES[slug]["nome"] for slug in especialidades_selecionadas]
    rotulo_especialidades = ", ".join(nomes_especialidades)

    # Monta a lista de tarefas a rodar em paralelo:
    # - Doctoralia e Sechat entram uma vez POR especialidade selecionada
    #   (cada um tem uma URL diferente por especialidade).
    # - As fontes genéricas (Ama-me, Kaya Doc, Cannaceia) não filtram por
    #   especialidade — a página é a mesma não importa o que você busca —
    #   então elas entram só UMA VEZ na lista toda.
    tarefas = []  # cada item: (rótulo_pra_erro, função, args)

    for slug in especialidades_selecionadas:
        info = ESPECIALIDADES[slug]
        nome_especialidade = info["nome"]

        tarefas.append((
            f"Doctoralia ({nome_especialidade})",
            buscar_medicos_doctoralia,
            (info["doctoralia"], nome_especialidade, cidade, uf),
        ))

        if info["sechat"]:
            tarefas.append((
                f"Sechat ({nome_especialidade})",
                buscar_medicos_sechat,
                (info["sechat"], nome_especialidade, cidade, uf),
            ))

    for chave_fonte, dados_fonte in FONTES_GENERICAS.items():
        tarefas.append((
            dados_fonte["nome"],
            buscar_medicos_fonte_generica,
            (chave_fonte, rotulo_especialidades, cidade, uf),
        ))

    todos_medicos = []
    fontes_com_erro = []

    # Roda tudo em paralelo (até 12 buscas ao mesmo tempo) em vez de uma
    # atrás da outra — é isso que evita o erro de timeout quando várias
    # especialidades são selecionadas juntas.
    with ThreadPoolExecutor(max_workers=12) as executor:
        futuros = {
            executor.submit(funcao, *args): rotulo
            for rotulo, funcao, args in tarefas
        }
        for futuro in as_completed(futuros):
            rotulo = futuros[futuro]
            try:
                todos_medicos.extend(futuro.result())
            except Exception as e:
                fontes_com_erro.append(f"{rotulo}: {e}")

    # Marca quais médicos já estão no painel do usuário logado, pra
    # mostrar "Presente no painel médico" em vez do botão de adicionar.
    usuario = usuario_logado()
    chaves_do_painel = chaves_no_painel_usuario(usuario["id"])
    for medico in todos_medicos:
        medico["no_painel"] = computar_chave_medico(medico) in chaves_do_painel

    return jsonify({
        "especialidades": nomes_especialidades,
        "cidade": cidade,
        "uf": uf,
        "total": len(todos_medicos),
        "medicos": todos_medicos,
        "avisos": fontes_com_erro,
    })


@app.route("/api/painel", methods=["GET"])
@login_required
def api_painel_listar():
    usuario = usuario_logado()
    painel = listar_painel_usuario(usuario["id"])
    return jsonify({
        "prospectados": [dict(m) for m in painel if m["status"] == "prospectado"],
        "visitados": [dict(m) for m in painel if m["status"] == "visitado"],
    })


@app.route("/api/painel/adicionar", methods=["POST"])
@login_required
def api_painel_adicionar():
    usuario = usuario_logado()
    medico = request.get_json(force=True) or {}

    if not medico.get("nome"):
        return jsonify({"ok": False, "mensagem": "Dados do médico incompletos."}), 400

    ok, mensagem = adicionar_ao_painel(usuario["id"], medico)
    return jsonify({"ok": ok, "mensagem": mensagem})


@app.route("/api/painel/mover", methods=["POST"])
@login_required
def api_painel_mover():
    usuario = usuario_logado()
    dados = request.get_json(force=True) or {}
    entrada_id = dados.get("id")
    novo_status = dados.get("status")

    if not entrada_id or novo_status not in ("prospectado", "visitado"):
        return jsonify({"ok": False, "mensagem": "Requisição inválida."}), 400

    ok = mover_no_painel(usuario["id"], int(entrada_id), novo_status)
    return jsonify({"ok": ok})


def abrir_navegador():
    webbrowser.open("http://127.0.0.1:5000")


if __name__ == "__main__":
    threading.Timer(1.2, abrir_navegador).start()
    app.run(debug=False, use_reloader=False, port=5000)
